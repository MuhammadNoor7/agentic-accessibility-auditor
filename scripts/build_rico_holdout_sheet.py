"""
Build import file for Google Sheet "RICO Dataset" tab from data/data-rico-holdout.
Pairs screenshots/{category}/{id}.jpg with xml/{category}/{id}.xml (1698 pairs).

Sheet columns (3 only): Category | Screenshot (embedded image) | Xml (hyperlink)
"""

import csv
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font
from PIL import Image

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = PROJECT_ROOT / "data" / "data-rico-holdout"
SCREENSHOT_ROOT = DATA_ROOT / "screenshots"
XML_ROOT = DATA_ROOT / "xml"
MANIFEST = DATA_ROOT / "manifest"
OUT_XLSX = MANIFEST / "rico_holdout_dataset_import.xlsx"
OUT_CSV = MANIFEST / "rico_holdout_dataset_import.csv"
OUT_CATEGORY_DIR = MANIFEST / "category_sheets"
OUT_GOOGLE_DIR = MANIFEST / "google_sheets"

THUMB_WIDTH = 120
THUMB_HEIGHT = 220
ROW_HEIGHT = 165
CATEGORY_COL_WIDTH = 14
IMG_COL_WIDTH = 18
XML_COL_WIDTH = 22
GOOGLE_SHEET_CHUNK_SIZE = 120

HYPERLINK_FONT = Font(color="0563C1", underline="single")

CATEGORY_DIRS = [
    "chat",
    "home",
    "list",
    "login",
    "maps",
    "menu",
    "profile",
    "search",
    "settings",
    "welcome",
]

DIR_TO_CLASS = {
    "chat": "Chat",
    "home": "Home",
    "list": "List",
    "login": "Login",
    "maps": "Map",
    "menu": "Menu",
    "profile": "Profile",
    "search": "Search",
    "settings": "Setting",
    "welcome": "Welcome",
}

CLASS_ORDER = [
    "Menu",
    "Login",
    "Home",
    "Setting",
    "Search",
    "List",
    "Map",
    "Chat",
    "Welcome",
    "Profile",
]


def screen_id_key(path: Path) -> int:
    return int(path.stem)


def xml_link(category: str, xml_name: str) -> str:
    return f"http://xml/{category}/{xml_name}"


def collect_pairs():
    pairs = []
    missing_xml = []

    for category in CATEGORY_DIRS:
        screen_class = DIR_TO_CLASS[category]
        screenshot_dir = SCREENSHOT_ROOT / category
        if not screenshot_dir.is_dir():
            continue

        for img_path in sorted(screenshot_dir.glob("*.jpg"), key=screen_id_key):
            xml_path = XML_ROOT / category / f"{img_path.stem}.xml"
            if not xml_path.is_file():
                missing_xml.append(str(xml_path))
                continue
            pairs.append((screen_class, category, img_path, xml_path))

    if missing_xml:
        raise FileNotFoundError(f"Missing {len(missing_xml)} XML files. Example: {missing_xml[0]}")

    class_rank = {name: index for index, name in enumerate(CLASS_ORDER)}
    pairs.sort(key=lambda item: (class_rank.get(item[0], 99), screen_id_key(item[2])))
    return pairs


def make_thumbnail(path: Path, out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(path) as img:
        img = img.convert("RGB")
        img.thumbnail((THUMB_WIDTH, THUMB_HEIGHT), Image.Resampling.LANCZOS)
        img.save(out_path, format="PNG")
    return out_path


def build_xlsx(pairs, out_path: Path, sheet_title: str, thumb_dir: Path) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws["A1"] = "Category"
    ws["B1"] = "Screenshot"
    ws["C1"] = "Xml"

    ws.column_dimensions["A"].width = CATEGORY_COL_WIDTH
    ws.column_dimensions["B"].width = IMG_COL_WIDTH
    ws.column_dimensions["C"].width = XML_COL_WIDTH

    for row_idx, (screen_class, category, img_path, xml_path) in enumerate(pairs, start=2):
        ws.row_dimensions[row_idx].height = ROW_HEIGHT
        ws.cell(row=row_idx, column=1, value=screen_class)

        thumb_path = thumb_dir / f"{category}_{img_path.stem}.png"
        make_thumbnail(img_path, thumb_path)
        xl_img = XLImage(str(thumb_path))
        xl_img.width = THUMB_WIDTH
        xl_img.height = THUMB_HEIGHT
        ws.add_image(xl_img, f"B{row_idx}")

        link_cell = ws.cell(row=row_idx, column=3, value=xml_path.name)
        link_cell.hyperlink = xml_link(category, xml_path.name)
        link_cell.font = HYPERLINK_FONT

    wb.save(out_path)
    return len(pairs)


def build_csv(pairs):
    with OUT_CSV.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Category", "Screenshot", "Xml", "Xml_Link"])
        for screen_class, category, img_path, xml_path in pairs:
            writer.writerow(
                [
                    screen_class,
                    img_path.name,
                    xml_path.name,
                    xml_link(category, xml_path.name),
                ]
            )


def build_category_sheets(pairs, thumb_dir: Path) -> None:
    OUT_CATEGORY_DIR.mkdir(parents=True, exist_ok=True)
    by_class = defaultdict(list)
    for item in pairs:
        by_class[item[0]].append(item)

    for screen_class in CLASS_ORDER:
        category_pairs = by_class.get(screen_class, [])
        if not category_pairs:
            continue
        out_path = OUT_CATEGORY_DIR / f"{screen_class}.xlsx"
        count = build_xlsx(category_pairs, out_path, screen_class, thumb_dir)
        size_mb = out_path.stat().st_size / 1024 / 1024
        print(f"  {screen_class:8s}: {count:3d} screenshots -> {out_path.name} ({size_mb:.1f} MB)")


def build_google_sheet_parts(pairs, thumb_dir: Path) -> None:
    OUT_GOOGLE_DIR.mkdir(parents=True, exist_ok=True)
    for old in OUT_GOOGLE_DIR.glob("rico_holdout_part_*.xlsx"):
        old.unlink()

    total_parts = math.ceil(len(pairs) / GOOGLE_SHEET_CHUNK_SIZE)
    for part_idx in range(total_parts):
        start = part_idx * GOOGLE_SHEET_CHUNK_SIZE
        chunk = pairs[start : start + GOOGLE_SHEET_CHUNK_SIZE]
        out_path = OUT_GOOGLE_DIR / f"rico_holdout_part_{part_idx + 1:02d}.xlsx"
        count = build_xlsx(chunk, out_path, "RICO Dataset", thumb_dir)
        size_mb = out_path.stat().st_size / 1024 / 1024
        print(f"  part {part_idx + 1:02d}/{total_parts:02d}: {count:3d} screenshots ({size_mb:.1f} MB)")


def main():
    if not SCREENSHOT_ROOT.is_dir() or not XML_ROOT.is_dir():
        raise SystemExit(f"Expected folders not found under {DATA_ROOT}")

    pairs = collect_pairs()
    print(f"Found {len(pairs)} screenshot/XML pairs.")

    MANIFEST.mkdir(parents=True, exist_ok=True)
    thumb_dir = MANIFEST / "_sheet_thumbs"
    thumb_dir.mkdir(parents=True, exist_ok=True)

    build_csv(pairs)
    print(f"Wrote {OUT_CSV}")

    print("\nBuilding per-category xlsx (import these — images show in Google Sheets):")
    build_category_sheets(pairs, thumb_dir)

    print(f"\nBuilding chunked xlsx for one RICO Dataset tab ({GOOGLE_SHEET_CHUNK_SIZE} rows each):")
    build_google_sheet_parts(pairs, thumb_dir)

    print("\nBuilding full combined xlsx (Excel only, images lost in Google Sheets):")
    count = build_xlsx(pairs, OUT_XLSX, "RICO Dataset", thumb_dir)
    print(f"  {count} screenshots -> {OUT_XLSX.name}")

    print("\nImport into Google Sheet (screenshots visible):")
    print("  Option A — one tab per category:")
    print("    Import category_sheets/Chat.xlsx into your Chat tab (same as MASC example)")
    print("  Option B — single RICO Dataset tab:")
    print("    Import google_sheets/rico_holdout_part_01.xlsx, then part_02 ... part_15")
    print("    Use 'Insert new rows' for parts 02-15")
    print("  Do NOT import rico_holdout_dataset_import.csv or the full 32 MB xlsx")


if __name__ == "__main__":
    main()
