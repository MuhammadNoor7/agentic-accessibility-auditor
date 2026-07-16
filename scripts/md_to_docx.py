"""Convert Markdown (.md) to a formatted Word document (.docx).

Usage:
  python scripts/md_to_docx.py docs/progress/Supplementary_Progress_Report_v1.0.md
  python scripts/md_to_docx.py input.md -o output.docx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
NUMBERED_RE = re.compile(r"^(\d+)\.\s+(.*)$")


def _set_cell_shading(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def _set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: Pt | None = None,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = size
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_formatted_run(paragraph, text: str, *, base_size: Pt = Pt(11)) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            _set_run_font(run, size=base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, size=base_size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, name="Consolas", size=Pt(9), color=RGBColor(0x1E, 0x3A, 0x5F))
        else:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, size=base_size)
            run.italic = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_run_font(run, size=base_size)


def _set_cell_text(cell, text: str, *, header: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    if header:
        run = paragraph.add_run(re.sub(r"[*`]", "", text))
        _set_run_font(run, size=Pt(10), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    else:
        _add_formatted_run(paragraph, text, base_size=Pt(10))


def _parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|?[\s:\-|]+\|?", line.strip()))


def _style_table(table) -> None:
    table.style = "Table Grid"
    table.autofit = True
    for cell in table.rows[0].cells:
        _set_cell_shading(cell, "1E3A5F")


def md_to_docx(md_path: Path, docx_path: Path) -> Path:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    md_dir = md_path.parent
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.95)
        section.right_margin = Inches(0.95)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    style.paragraph_format.space_after = Pt(6)

    for level in range(1, 5):
        heading = doc.styles[f"Heading {level}"]
        heading.font.name = "Calibri"
        heading.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
        heading.font.bold = True
        sizes = {1: 18, 2: 14, 3: 12, 4: 11}
        heading.font.size = Pt(sizes[level])

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1 and i < 5:
                p = doc.add_heading(text, level=1)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_heading(text, level=min(level, 4))
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and _is_table_separator(lines[i + 1]):
            headers = _parse_table_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_parse_table_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(headers))
            _style_table(table)
            for col, header in enumerate(headers):
                _set_cell_text(table.rows[0].cells[col], header, header=True)
            for row_idx, row in enumerate(rows, start=1):
                for col, cell in enumerate(row):
                    if col < len(table.rows[row_idx].cells):
                        _set_cell_text(table.rows[row_idx].cells[col], cell, header=False)
            doc.add_paragraph()
            continue
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            run_prefix = paragraph.add_run("| ")
            _set_run_font(run_prefix, bold=True, color=RGBColor(0x1E, 0x3A, 0x5F))
            _add_formatted_run(paragraph, stripped.lstrip("> ").strip())
            i += 1
            continue
        if stripped.startswith("- ") or stripped.startswith("* "):
            paragraph = doc.add_paragraph(style="List Bullet")
            _add_formatted_run(paragraph, stripped[2:].strip())
            i += 1
            continue
        numbered = NUMBERED_RE.match(stripped)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            _add_formatted_run(paragraph, numbered.group(2))
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.15)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run("\n".join(code_lines))
            _set_run_font(run, name="Consolas", size=Pt(8.5), color=RGBColor(0x1E, 0x3A, 0x5F))
            i += 1
            continue
        image_match = IMAGE_RE.match(stripped)
        if image_match:
            alt, rel_path = image_match.group(1), image_match.group(2)
            image_path = (md_dir / rel_path).resolve()
            if alt:
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(alt)
                run.italic = True
                _set_run_font(run, size=Pt(9), color=RGBColor(0x55, 0x65, 0x7A))
            if image_path.is_file():
                doc.add_picture(str(image_path), width=Inches(5.8))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                missing = doc.add_paragraph()
                missing.add_run(f"[Missing image: {rel_path}]")
            i += 1
            continue

        paragraph = doc.add_paragraph()
        _add_formatted_run(paragraph, stripped)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    return docx_path


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert Markdown to DOCX")
    parser.add_argument("input", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    args = parser.parse_args()
    output = args.output or args.input.with_suffix(".docx")
    result = md_to_docx(args.input.resolve(), output.resolve())
    print(f"Wrote {result}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
